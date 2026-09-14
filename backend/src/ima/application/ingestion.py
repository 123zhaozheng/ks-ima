"""Bounded parsing and structure-aware chunking for approved ingestion formats."""

# ruff: noqa: E501

from __future__ import annotations

import json
import math
import re
from collections.abc import Callable, Mapping
from dataclasses import dataclass, field, replace
from hashlib import sha256
from io import BytesIO
from pathlib import PurePosixPath
from typing import cast
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader

from ima.config import Settings

SUPPORTED_MIME_TYPES = frozenset(
    {
        "text/plain",
        "text/markdown",
        "application/json",
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
)

CHUNKER_VERSION = "structure-aware-v1"
_TOKEN_ESTIMATOR_VERSION = "cjk-1-latin-quarter-v1"
_CJK_RE = re.compile(
    r"[\u2e80-\u2fff\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uac00-\ud7af]"
)
_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_SETEXT_RE = re.compile(r"^\s*(=+|-+)\s*$")
_CODE_FENCE_RE = re.compile(r"^\s*(`{3,}|~{3,})(.*)$")
_LIST_RE = re.compile(r"^\s*(?:[-*+]\s+|\d+[.)]\s+)")
_TABLE_SEPARATOR_RE = re.compile(r"^\s*:?-{3,}:?\s*$")
_SENTENCE_SPLIT_RE = re.compile(r"(?<=[。！？!?；;])\s*|\n+|(?<=[A-Za-z]\.)\s+(?=[\"'“‘A-Z0-9])")


class ParseError(Exception):
    def __init__(self, code: str, *, retryable: bool = False) -> None:
        super().__init__(code)
        self.code = code
        self.retryable = retryable


@dataclass(frozen=True, slots=True)
class ParsedBlock:
    """A durable, JSON-compatible unit emitted by a format parser."""

    kind: str
    text: str
    heading_path: tuple[str, ...] = ()
    heading_level: int | None = None
    page_start: int | None = None
    page_end: int | None = None
    sheet_name: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    metadata: dict[str, object] = field(default_factory=dict)
    has_header: bool | None = None


@dataclass(frozen=True, slots=True)
class ParsedText:
    parser: str
    parser_version: str
    text: str
    digest: str
    blocks: tuple[ParsedBlock, ...] = ()

    @property
    def ir(self) -> dict[str, object]:
        """Return the intermediate representation ready for a JSONB column."""
        return {
            "version": 1,
            "blocks": [block_to_dict(block) for block in self.blocks],
        }


@dataclass(frozen=True, slots=True)
class ChunkRecord:
    """A chunk and the citation metadata generated alongside its text."""

    ordinal: int
    text_content: str
    content_digest: str
    heading_path: tuple[str, ...]
    page_start: int | None
    page_end: int | None
    sheet_name: str | None
    chunk_type: str
    token_count: int
    chunker_version: str
    metadata: dict[str, object]

    # Keep the old three-value unpacking convenient for callers that only need
    # the text corpus. New persistence code uses the named fields above.
    def as_legacy_tuple(self) -> tuple[int, str, str]:
        return self.ordinal, self.text_content, self.content_digest


def _json_value(value: object) -> object:
    if isinstance(value, tuple):
        return [_json_value(item) for item in value]
    if isinstance(value, list):
        return [_json_value(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    return value


def block_to_dict(block: ParsedBlock) -> dict[str, object]:
    result: dict[str, object] = {
        "kind": block.kind,
        "text": block.text,
        "heading_path": list(block.heading_path),
        "heading_level": block.heading_level,
        "page_start": block.page_start,
        "page_end": block.page_end,
        "sheet_name": block.sheet_name,
        "row_start": block.row_start,
        "row_end": block.row_end,
        "has_header": block.has_header,
        "metadata": _json_value(block.metadata),
    }
    return result


def blocks_from_ir(value: object, fallback_text: str = "") -> tuple[ParsedBlock, ...]:
    """Decode persisted IR and fall back safely for rows created pre-IR."""
    raw_blocks: object = value
    if isinstance(value, Mapping):
        raw_blocks = value.get("blocks", ())
    if not isinstance(raw_blocks, list | tuple):
        return text_blocks(fallback_text)

    blocks: list[ParsedBlock] = []
    for raw in raw_blocks:
        if not isinstance(raw, Mapping):
            continue
        kind = raw.get("kind")
        text = raw.get("text")
        if not isinstance(kind, str) or not isinstance(text, str) or not text.strip():
            continue
        raw_path = raw.get("heading_path", ())
        path = tuple(str(item) for item in raw_path) if isinstance(raw_path, list | tuple) else ()
        raw_metadata = raw.get("metadata", {})
        metadata = dict(raw_metadata) if isinstance(raw_metadata, Mapping) else {}
        blocks.append(
            ParsedBlock(
                kind=kind,
                text=text,
                heading_path=path,
                heading_level=_optional_int(raw.get("heading_level")),
                page_start=_optional_int(raw.get("page_start")),
                page_end=_optional_int(raw.get("page_end")),
                sheet_name=_optional_str(raw.get("sheet_name")),
                row_start=_optional_int(raw.get("row_start")),
                row_end=_optional_int(raw.get("row_end")),
                has_header=raw.get("has_header")
                if isinstance(raw.get("has_header"), bool)
                else None,
                metadata=cast(dict[str, object], metadata),
            )
        )
    return tuple(blocks) or text_blocks(fallback_text)


def _optional_int(value: object) -> int | None:
    return value if isinstance(value, int) and not isinstance(value, bool) else None


def _optional_str(value: object) -> str | None:
    return value if isinstance(value, str) else None


def validate_upload_type(filename: str, mime_type: str) -> None:
    suffix = PurePosixPath(filename).suffix.lower()
    allowed = {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".json": "application/json",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    if mime_type.lower() not in SUPPORTED_MIME_TYPES or allowed.get(suffix) != mime_type.lower():
        raise ParseError("UNSUPPORTED_FILE_TYPE")


class Parser:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def parse(self, data: bytes, filename: str, mime_type: str) -> ParsedText:
        if len(data) > self.settings.storage_max_object_bytes:
            raise ParseError("OBJECT_TOO_LARGE")
        validate_upload_type(filename, mime_type)
        suffix = PurePosixPath(filename).suffix.lower()
        try:
            if suffix in {".docx", ".xlsx"}:
                self._validate_archive(data)
            if suffix in {".txt", ".md"}:
                parser = "markdown" if suffix == ".md" else "text"
                blocks = text_blocks(data.decode("utf-8"))
            elif suffix == ".json":
                parser = "json"
                value = json.dumps(json.loads(data.decode("utf-8")), ensure_ascii=False, indent=2)
                blocks = (ParsedBlock("prose", value),)
            elif suffix == ".pdf":
                parser, blocks = "pdf", self._pdf(data)
            elif suffix == ".docx":
                parser, blocks = "docx", self._docx(data)
            elif suffix == ".xlsx":
                parser, blocks = "spreadsheet", self._spreadsheet(data)
            else:
                raise ParseError("UNSUPPORTED_FILE_TYPE")
        except ParseError:
            raise
        except (OSError, UnicodeDecodeError, ValueError, KeyError, TypeError) as exc:
            raise ParseError("MALFORMED_DOCUMENT") from exc
        blocks = _normalized_blocks(blocks)
        normalized = _normalize("\n\n".join(block.text for block in blocks))
        if not normalized:
            raise ParseError("EMPTY_DOCUMENT")
        encoded = normalized.encode("utf-8")
        if len(encoded) > self.settings.ingestion_max_text_bytes:
            raise ParseError("EXTRACTED_TEXT_TOO_LARGE")
        return ParsedText(parser, "2", normalized, sha256(encoded).hexdigest(), blocks)

    def _validate_archive(self, data: bytes) -> None:
        try:
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000:
                    raise ParseError("ARCHIVE_ENTRY_LIMIT")
                compressed = sum(entry.compress_size for entry in entries)
                unpacked = sum(entry.file_size for entry in entries)
                if unpacked > self.settings.ingestion_max_text_bytes * 20 or (
                    compressed and unpacked / compressed > 100
                ):
                    raise ParseError("ARCHIVE_LIMIT")
        except BadZipFile as exc:
            raise ParseError("MALFORMED_DOCUMENT") from exc

    def _pdf(self, data: bytes) -> tuple[ParsedBlock, ...]:
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ParseError("ENCRYPTED_DOCUMENT")
        if len(reader.pages) > 1000:
            raise ParseError("PDF_PAGE_LIMIT")
        blocks: list[ParsedBlock] = []
        for page_number, page in enumerate(reader.pages, start=1):
            value = page.extract_text() or ""
            if value.strip():
                blocks.append(
                    ParsedBlock(
                        "prose",
                        value,
                        page_start=page_number,
                        page_end=page_number,
                        metadata={"page": page_number},
                    )
                )
        return tuple(blocks)

    def _docx(self, data: bytes) -> tuple[ParsedBlock, ...]:
        document = Document(BytesIO(data))
        blocks: list[ParsedBlock] = []
        heading_stack: list[str] = []
        body = document.element.body
        for child in body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                value = paragraph.text.strip()
                if not value:
                    continue
                level = _heading_level(paragraph.style.name if paragraph.style else "")
                if level is not None:
                    heading_stack = _set_heading(heading_stack, level, value)
                    blocks.append(
                        ParsedBlock("heading", value, tuple(heading_stack), heading_level=level)
                    )
                else:
                    kind = (
                        "list"
                        if paragraph.style and "list" in paragraph.style.name.lower()
                        else "prose"
                    )
                    blocks.append(ParsedBlock(kind, value, tuple(heading_stack)))
            elif child.tag.endswith("}tbl"):
                table = Table(child, document)
                rows = [
                    _format_cells([cell.text.strip() for cell in row.cells]) for row in table.rows
                ]
                rows = [row for row in rows if row]
                if rows:
                    has_header = _docx_table_has_header(table)
                    header = rows[0] if has_header else None
                    blocks.append(
                        ParsedBlock(
                            "table",
                            "\n".join(rows),
                            tuple(heading_stack),
                            has_header=has_header,
                            metadata={
                                "rows": rows,
                                "has_header": has_header,
                                "header": header,
                                "table_header": header,
                            },
                            row_start=1,
                            row_end=len(rows),
                        )
                    )
        return tuple(blocks)

    def _spreadsheet(self, data: bytes) -> tuple[ParsedBlock, ...]:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        if len(workbook.sheetnames) > 100:
            raise ParseError("SHEET_LIMIT")
        blocks: list[ParsedBlock] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            raw_rows: list[list[object]] = []
            bold_rows: list[bool] = []
            first_row: int | None = None
            last_row = 0
            for row_number, row in enumerate(sheet.iter_rows(values_only=False), start=1):
                if row_number > 100_000:
                    raise ParseError("SHEET_LIMIT")
                raw_values = [cell.value for cell in row]
                values = ["" if value is None else str(value) for value in raw_values]
                if not any(value.strip() for value in values):
                    continue
                if first_row is None:
                    first_row = row_number
                raw_rows.append(raw_values)
                bold_rows.append(
                    any(
                        bool(getattr(getattr(cell, "font", None), "bold", False))
                        for cell in row
                        if cell.value is not None
                    )
                )
                rows.append(_format_cells(values))
                last_row = row_number
            if not rows:
                continue
            has_header = _xlsx_has_header(raw_rows, bold_rows)
            header = rows[0] if has_header else None
            blocks.append(
                ParsedBlock(
                    "heading",
                    sheet.title,
                    (sheet.title,),
                    heading_level=1,
                    sheet_name=sheet.title,
                    metadata={"sheet_name": sheet.title},
                )
            )
            blocks.append(
                ParsedBlock(
                    "table",
                    "\n".join(rows),
                    (sheet.title,),
                    sheet_name=sheet.title,
                    row_start=first_row,
                    row_end=last_row,
                    has_header=has_header,
                    metadata={
                        "rows": rows,
                        "has_header": has_header,
                        "header": header,
                        "table_header": header,
                        "sheet_name": sheet.title,
                        "row_start": first_row,
                        "row_end": last_row,
                    },
                )
            )
        return tuple(blocks)

    @staticmethod
    def _normalize(value: str) -> str:
        return _normalize(value)


def _normalize(value: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", value.replace("\r\n", "\n").replace("\r", "\n")).strip()


def _normalized_blocks(blocks: tuple[ParsedBlock, ...]) -> tuple[ParsedBlock, ...]:
    normalized: list[ParsedBlock] = []
    for block in blocks:
        value = _normalize(block.text)
        if value:
            normalized.append(replace(block, text=value))
    return tuple(normalized)


def _heading_level(style_name: str) -> int | None:
    match = re.search(r"(?:heading|标题)\s*([1-6])$", style_name.strip(), re.IGNORECASE)
    return int(match.group(1)) if match else None


def _set_heading(stack: list[str], level: int, title: str) -> list[str]:
    return [*stack[: max(0, level - 1)], title]


def _is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return len(cells) >= 1 and all(_TABLE_SEPARATOR_RE.fullmatch(cell) for cell in cells)


def _format_cells(values: list[str]) -> str:
    return " | ".join(value.strip() for value in values).strip()


def _format_table_line(line: str) -> str:
    return _format_cells([cell for cell in line.strip().strip("|").split("|")])


def _xlsx_has_header(rows: list[list[object]], bold_rows: list[bool]) -> bool:
    """Infer an XLSX header from read-only cell types or visible styling.

    ``ReadOnlyWorksheet`` deliberately omits ``auto_filter``. A string-only
    first row followed by a typed value is a dependable header signal; bold
    formatting is the safe fallback when the data rows are also strings.
    """
    if len(rows) < 2:
        return False
    first_values = [value for value in rows[0] if value is not None and str(value).strip()]
    if not first_values or not all(isinstance(value, str) for value in first_values):
        return False
    for data_row in rows[1:]:
        if any(value is not None and not isinstance(value, str) for value in data_row):
            return True
    return bool(bold_rows and bold_rows[0] and not any(bold_rows[1:]))


def _docx_table_has_header(table: Table) -> bool:
    """Honor the OOXML repeat-header marker instead of guessing from row text."""
    if not table.rows:
        return False
    properties = table.rows[0]._tr.find(qn("w:trPr"))
    return properties is not None and properties.find(qn("w:tblHeader")) is not None


def _looks_like_table(lines: list[str], index: int) -> bool:
    if "|" not in lines[index]:
        return False
    if index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
        return True
    return index + 1 < len(lines) and "|" in lines[index + 1]


def text_blocks(value: str) -> tuple[ParsedBlock, ...]:
    """Parse Markdown/text into headings, prose, lists, code, and tables."""
    lines = value.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[ParsedBlock] = []
    heading_stack: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_stack = _set_heading(heading_stack, level, title)
            blocks.append(ParsedBlock("heading", title, tuple(heading_stack), heading_level=level))
            index += 1
            continue
        if (
            index + 1 < len(lines)
            and lines[index + 1].strip()
            and _SETEXT_RE.fullmatch(lines[index + 1])
            and line.strip()
        ):
            level = 1 if lines[index + 1].strip().startswith("=") else 2
            title = line.strip()
            heading_stack = _set_heading(heading_stack, level, title)
            blocks.append(ParsedBlock("heading", title, tuple(heading_stack), heading_level=level))
            index += 2
            continue

        fence = _CODE_FENCE_RE.match(line)
        if fence:
            marker = fence.group(1)[0]
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].lstrip().startswith(
                marker * len(fence.group(1))
            ):
                code_lines.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            blocks.append(
                ParsedBlock(
                    "code",
                    "\n".join(code_lines),
                    tuple(heading_stack),
                    metadata={"language": fence.group(2).strip()},
                )
            )
            continue

        if _looks_like_table(lines, index):
            table_lines: list[str] = []
            has_header = index + 1 < len(lines) and _is_table_separator(lines[index + 1])
            while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                if not _is_table_separator(lines[index]):
                    table_lines.append(_format_table_line(lines[index]))
                index += 1
            if table_lines:
                blocks.append(
                    ParsedBlock(
                        "table",
                        "\n".join(table_lines),
                        tuple(heading_stack),
                        has_header=has_header,
                        metadata={
                            "rows": table_lines,
                            "has_header": has_header,
                            "header": table_lines[0] if has_header else None,
                            "table_header": table_lines[0] if has_header else None,
                        },
                        row_start=1,
                        row_end=len(table_lines),
                    )
                )
            continue

        if _LIST_RE.match(line):
            list_lines: list[str] = []
            while index < len(lines) and lines[index].strip() and _LIST_RE.match(lines[index]):
                list_lines.append(lines[index].strip())
                index += 1
            blocks.append(ParsedBlock("list", "\n".join(list_lines), tuple(heading_stack)))
            continue

        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines) and lines[index].strip():
            if (
                _HEADING_RE.match(lines[index])
                or _CODE_FENCE_RE.match(lines[index])
                or _LIST_RE.match(lines[index])
                or _looks_like_table(lines, index)
                or (
                    index + 1 < len(lines)
                    and _SETEXT_RE.fullmatch(lines[index + 1])
                    and lines[index + 1].strip()
                )
            ):
                break
            paragraph_lines.append(lines[index].strip())
            index += 1
        blocks.append(ParsedBlock("prose", "\n".join(paragraph_lines), tuple(heading_stack)))
    return tuple(blocks)


def estimate_tokens(text: str) -> int:
    """Estimate mixed-language tokens: CJK is ~1 token/char, Latin ~1/4 char."""
    cjk = sum(1 for character in text if _CJK_RE.fullmatch(character))
    latin_and_other = len(text) - cjk
    return cjk + math.ceil(latin_and_other / 4)


def chunker_config_digest(
    size: int,
    overlap_blocks: int,
    maximum: int,
    chunker_version: str = CHUNKER_VERSION,
) -> str:
    config = {
        "chunker_version": chunker_version,
        "token_budget": size,
        "overlap_blocks": overlap_blocks,
        "max_chunks": maximum,
        "token_estimator": _TOKEN_ESTIMATOR_VERSION,
    }
    encoded = json.dumps(config, ensure_ascii=True, separators=(",", ":"), sort_keys=True)
    return sha256(encoded.encode("utf-8")).hexdigest()


@dataclass(frozen=True, slots=True)
class _Piece:
    text: str
    source: ParsedBlock
    metadata: dict[str, object] = field(default_factory=dict)


def _prefix(path: tuple[str, ...]) -> str:
    return " > ".join(item for item in path if item.strip())


def _fit_heading_path(path: tuple[str, ...], budget: int) -> tuple[str, ...]:
    """Keep a useful heading suffix when the full path consumes the budget."""
    prefix = _prefix(path)
    available = budget - estimate_tokens("\n\n") - 1
    if not prefix or estimate_tokens(prefix) <= available:
        return path
    if available < 1:
        return ()

    suffix = ""
    for item in reversed(path):
        candidate = item if not suffix else f"{item} > {suffix}"
        if estimate_tokens(candidate) <= available:
            suffix = candidate
            continue
        if not suffix:
            characters = ""
            for unit in reversed(_fallback_units(item)):
                candidate_text = unit + characters
                if estimate_tokens(candidate_text) > available:
                    break
                characters = candidate_text
            suffix = characters.strip()
        break
    return (suffix,) if suffix else ()


def _render(pieces: list[_Piece], path: tuple[str, ...]) -> str:
    body = "\n\n".join(piece.text.strip() for piece in pieces if piece.text.strip())
    heading = _prefix(path)
    return f"{heading}\n\n{body}" if heading and body else heading or body


def _body_budget(path: tuple[str, ...], budget: int) -> int:
    # ``_render`` inserts a separator between the heading prefix and body.
    return max(1, budget - estimate_tokens(_prefix(path)) - estimate_tokens("\n\n"))


def _split_by_char_budget(value: str, budget: int) -> list[str]:
    result: list[str] = []
    current = ""
    for unit in _fallback_units(value):
        candidate = current + unit
        if current and estimate_tokens(candidate) > budget:
            result.append(current.strip())
            current = ""
            candidate = unit
        if not current and estimate_tokens(candidate) > budget:
            # Keep a known abbreviation atomic. A one-token budget cannot
            # represent a multi-token abbreviation without breaking it.
            result.append(unit.strip())
            continue
        current = candidate
    if current.strip():
        result.append(current.strip())
    return [part for part in result if part]


def _split_text(value: str, budget: int, *, sentences: bool) -> tuple[str, ...]:
    if estimate_tokens(value) <= budget:
        return (value.strip(),)
    pieces = list(_sentence_pieces(value)) if sentences else [value.strip()]
    grouped: list[str] = []
    current = ""
    for piece in pieces:
        candidate = f"{current}\n{piece}" if current else piece
        if current and estimate_tokens(candidate) > budget:
            grouped.append(current.strip())
            current = ""
        if estimate_tokens(piece) > budget:
            if current:
                grouped.append(current.strip())
                current = ""
            grouped.extend(_split_by_char_budget(piece, budget))
        else:
            current = piece if not current else f"{current}\n{piece}"
    if current:
        grouped.append(current.strip())
    return tuple(piece for piece in grouped if piece)


def _split_to_render_budget(
    value: str, budget: int, render: Callable[[str], str]
) -> tuple[str, ...]:
    if estimate_tokens(render(value)) <= budget:
        return (value.strip(),)
    result: list[str] = []
    current = ""
    for unit in _fallback_units(value):
        candidate = current + unit
        if current and estimate_tokens(render(candidate)) > budget:
            result.append(current.strip())
            current = ""
            candidate = unit
        if not current and estimate_tokens(render(candidate)) > budget:
            # The heading path is already truncated before this function is
            # called. Keep the atomic abbreviation rather than returning no
            # chunks or cutting it in the middle.
            result.append(unit.strip())
            continue
        current = candidate
    if current.strip():
        result.append(current.strip())
    return tuple(result)


_ABBREVIATIONS = frozenset(
    {
        "mr.",
        "mrs.",
        "ms.",
        "dr.",
        "prof.",
        "sr.",
        "jr.",
        "st.",
        "inc.",
        "ltd.",
        "co.",
        "corp.",
        "vs.",
        "etc.",
        "e.g.",
        "i.e.",
        "fig.",
        "no.",
        "eq.",
        "ph.d.",
        "u.s.",
    }
)

_ABBREVIATION_RE = re.compile(
    r"(?<![A-Za-z])(?:"
    + "|".join(
        re.escape(abbreviation) for abbreviation in sorted(_ABBREVIATIONS, key=len, reverse=True)
    )
    + r")(?![A-Za-z])",
    re.IGNORECASE,
)


def _fallback_units(value: str) -> tuple[str, ...]:
    """Return character units while keeping known abbreviations indivisible."""
    units: list[str] = []
    cursor = 0
    for match in _ABBREVIATION_RE.finditer(value):
        units.extend(value[cursor : match.start()])
        units.append(match.group(0))
        cursor = match.end()
    units.extend(value[cursor:])
    return tuple(units)


def _safe_english_period(value: str, period: int) -> bool:
    """Reject decimal points, abbreviations, and URL dots as sentence ends."""
    before = value[:period]
    after = value[period + 1 :]
    if not before or not after:
        return True
    if before[-1].isdigit() or after[0].isdigit():
        return False
    word_match = re.search(r"[A-Za-z](?:[A-Za-z.]*)$", before)
    word = word_match.group(0).lower() if word_match else ""
    if f"{word}." in _ABBREVIATIONS or len(word) <= 2:
        return False
    token_start = max(before.rfind(" "), before.rfind("\n"), before.rfind("\t")) + 1
    token_end = len(after)
    for index, character in enumerate(after):
        if character.isspace():
            token_end = index
            break
    token_core = (before[token_start:] + "." + after[:token_end]).lower().rstrip(".,!?;:")
    if (
        "://" in token_core
        or token_core.startswith("www.")
        or re.search(r"[a-z0-9-]+\.[a-z]{2,}(?:[/:?#]|$)", token_core)
    ):
        return False
    return True


def _sentence_pieces(value: str) -> tuple[str, ...]:
    pieces: list[str] = []
    start = 0
    for match in _SENTENCE_SPLIT_RE.finditer(value):
        boundary = match.start()
        if (
            boundary
            and value[boundary - 1] == "."
            and not _safe_english_period(value, boundary - 1)
        ):
            continue
        piece = value[start:boundary].strip()
        if piece:
            pieces.append(piece)
        start = match.end()
    tail = value[start:].strip()
    if tail:
        pieces.append(tail)
    return tuple(pieces)


def _table_has_header(block: ParsedBlock) -> bool:
    if block.has_header is not None:
        return block.has_header
    value = block.metadata.get("has_header")
    return value if isinstance(value, bool) else False


def _table_pieces(block: ParsedBlock, budget: int, path: tuple[str, ...]) -> tuple[_Piece, ...]:
    raw_rows = block.metadata.get("rows", ())
    rows = (
        [str(row) for row in raw_rows]
        if isinstance(raw_rows, list | tuple)
        else block.text.splitlines()
    )
    rows = [row.strip() for row in rows if row.strip()]
    if not rows:
        return ()
    has_header = _table_has_header(block)
    header = rows[0] if has_header else None
    data_rows = rows[1:] if has_header else rows
    if not data_rows:
        return ()

    def make_piece(group: list[str], metadata: dict[str, object] | None = None) -> _Piece:
        return _Piece("\n".join(group), block, metadata or {})

    def fits(group: list[str]) -> bool:
        return estimate_tokens(_render([make_piece(group)], path)) <= budget

    if fits(rows):
        return (_Piece("\n".join(rows), block),)

    groups: list[list[str]] = []
    fallback: list[bool] = []
    current: list[str] = []
    for row in data_rows:
        candidate = [*([header] if header is not None else []), *current, row]
        if fits(candidate):
            current.append(row)
            continue
        if current:
            groups.append([*([header] if header is not None else []), *current])
            fallback.append(False)
            current = []
        candidate = [*([header] if header is not None else []), row]
        if fits(candidate):
            current = [row]
            continue

        row_parts = _split_to_render_budget(
            row,
            budget,
            lambda value: _render(
                [make_piece([*([header] if header is not None else []), value])], path
            ),
        )
        for row_part in row_parts:
            groups.append([*([header] if header is not None else []), row_part])
            fallback.append(True)
    if current:
        groups.append([*([header] if header is not None else []), *current])
        fallback.append(False)

    split = len(groups) > 1
    return tuple(
        make_piece(
            group,
            {
                **({"table_header_repeated": True} if has_header and split else {}),
                **({"table_split": True} if split else {}),
                **({"oversized_fallback": True} if was_fallback else {}),
            },
        )
        for group, was_fallback in zip(groups, fallback, strict=True)
    )


def _piece_for_block(block: ParsedBlock, budget: int, path: tuple[str, ...]) -> tuple[_Piece, ...]:
    if block.kind == "table":
        return _table_pieces(block, budget, path)

    def render(value: str) -> str:
        return _render([_Piece(value, block)], path)

    body_budget = _body_budget(path, budget)
    if estimate_tokens(render(block.text)) <= budget:
        return (_Piece(block.text, block),)
    values = _split_text(block.text, body_budget, sentences=block.kind in {"prose", "list"})
    result: list[_Piece] = []
    for value in values:
        for part in _split_to_render_budget(value, budget, render):
            result.append(_Piece(part, block, {"oversized_fallback": True, "atomic_split": True}))
    return tuple(result)


def _record(
    ordinal: int,
    pieces: list[_Piece],
    chunker_version: str,
    config_digest: str,
) -> ChunkRecord:
    path = next((piece.source.heading_path for piece in pieces if piece.source.heading_path), ())
    rendered = _render(pieces, path)
    pages = [
        page
        for piece in pieces
        for page in (piece.source.page_start, piece.source.page_end)
        if page is not None
    ]
    sheets = {piece.source.sheet_name for piece in pieces if piece.source.sheet_name}
    kinds = {piece.source.kind for piece in pieces}
    if kinds == {"table"}:
        chunk_type = "table"
    elif kinds == {"code"}:
        chunk_type = "code"
    elif not rendered or kinds == {"heading"}:
        chunk_type = "heading"
    else:
        chunk_type = "prose"
    metadata: dict[str, object] = {
        "chunker_config_digest": config_digest,
        "token_estimator": _TOKEN_ESTIMATOR_VERSION,
        "source_block_count": len(pieces),
    }
    for piece in pieces:
        metadata.update(piece.metadata)
    if pages:
        metadata["page_range"] = [min(pages), max(pages)]
    row_starts = [piece.source.row_start for piece in pieces if piece.source.row_start is not None]
    row_ends = [piece.source.row_end for piece in pieces if piece.source.row_end is not None]
    if row_starts and row_ends:
        metadata["row_range"] = [min(row_starts), max(row_ends)]
    digest = sha256(rendered.encode("utf-8")).hexdigest()
    return ChunkRecord(
        ordinal=ordinal,
        text_content=rendered,
        content_digest=digest,
        heading_path=path,
        page_start=min(pages) if pages else None,
        page_end=max(pages) if pages else None,
        sheet_name=next(iter(sheets)) if len(sheets) == 1 else None,
        chunk_type=chunk_type,
        token_count=estimate_tokens(rendered),
        chunker_version=chunker_version,
        metadata=metadata,
    )


def structure_aware_chunks(
    blocks: tuple[ParsedBlock, ...] | list[ParsedBlock],
    size: int,
    overlap_blocks: int,
    maximum: int,
    *,
    chunker_version: str = CHUNKER_VERSION,
    config_digest: str | None = None,
) -> tuple[ChunkRecord, ...]:
    """Chunk IR at headings/blocks first, then sentences/chars for oversized blocks."""
    if size < 1 or overlap_blocks < 0 or maximum < 1:
        raise ValueError("chunker limits must be positive and overlap_blocks cannot be negative")
    digest = config_digest or chunker_config_digest(size, overlap_blocks, maximum, chunker_version)
    records: list[ChunkRecord] = []
    pending: list[_Piece] = []
    overlap: list[_Piece] = []
    last_heading: ParsedBlock | None = None

    def emit(pieces: list[_Piece]) -> None:
        if not pieces:
            return
        records.append(_record(len(records), pieces, chunker_version, digest))
        if len(records) > maximum:
            raise ParseError("CHUNK_LIMIT")

    def flush(*, keep_overlap: bool = True) -> None:
        nonlocal pending, overlap
        if pending:
            emit(pending)
            overlap = pending[-overlap_blocks:] if keep_overlap and overlap_blocks else []
            pending = []

    for block in blocks:
        if not block.text.strip() and block.kind != "heading":
            continue
        if block.kind == "heading":
            flush(keep_overlap=False)
            overlap = []
            last_heading = block
            continue

        raw_path = block.heading_path or (last_heading.heading_path if last_heading else ())
        path = _fit_heading_path(raw_path, size)
        if block.heading_path != path:
            metadata = dict(block.metadata)
            if raw_path != path:
                metadata["full_heading_path"] = list(raw_path)
                metadata["heading_path_truncated"] = True
            block = replace(block, heading_path=path, metadata=metadata)
        pieces = _piece_for_block(block, size, path)
        if not pieces:
            continue
        if block.kind == "table" or len(pieces) > 1:
            flush(keep_overlap=False)
            for piece in pieces:
                emit([piece])
            overlap = []
            continue

        piece = pieces[0]
        if overlap:
            candidate = [*overlap, piece]
            if estimate_tokens(_render(candidate, path)) <= size:
                pending = candidate
            else:
                pending = [piece]
            overlap = []
            continue
        if pending and estimate_tokens(_render([*pending, piece], path)) > size:
            flush()
            if overlap:
                candidate = [*overlap, piece]
                pending = (
                    candidate if estimate_tokens(_render(candidate, path)) <= size else [piece]
                )
                overlap = []
            else:
                pending = [piece]
        else:
            pending.append(piece)

    flush(keep_overlap=False)
    return tuple(records)


def deterministic_chunks(
    text: str, size: int, overlap: int, maximum: int
) -> tuple[tuple[int, str, str], ...]:
    """Compatibility wrapper returning the historic three-value chunk tuples."""
    records = structure_aware_chunks(text_blocks(text), size, max(0, overlap), maximum)
    return tuple(record.as_legacy_tuple() for record in records)


__all__ = [
    "CHUNKER_VERSION",
    "ChunkRecord",
    "ParseError",
    "ParsedBlock",
    "ParsedText",
    "Parser",
    "blocks_from_ir",
    "chunker_config_digest",
    "deterministic_chunks",
    "estimate_tokens",
    "structure_aware_chunks",
    "text_blocks",
    "validate_upload_type",
]
