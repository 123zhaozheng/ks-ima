from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from ima.application.ingestion import (
    ParsedBlock,
    ParseError,
    Parser,
    deterministic_chunks,
    estimate_tokens,
    structure_aware_chunks,
    validate_upload_type,
)
from ima.config import Settings
from ima.infrastructure.storage import ObjectStorageClient


def test_parser_accepts_bounded_json_and_is_deterministic() -> None:
    parser = Parser(Settings())
    parsed = parser.parse(b'{"name":"IMA","items":[1,2]}', "sample.json", "application/json")
    assert parsed.parser == "json"
    assert parsed.digest == sha256(parsed.text.encode()).hexdigest()
    assert deterministic_chunks(parsed.text, 10, 2, 10) == deterministic_chunks(
        parsed.text, 10, 2, 10
    )


@pytest.mark.parametrize(
    ("filename", "mime_type"),
    [
        (
            "presentation.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
        ("image.png", "image/png"),
        ("archive.zip", "application/zip"),
    ],
)
def test_parser_rejects_out_of_scope_types(filename: str, mime_type: str) -> None:
    with pytest.raises(ParseError, match="UNSUPPORTED_FILE_TYPE"):
        validate_upload_type(filename, mime_type)


def test_parser_rejects_legacy_xls_without_a_biff_parser() -> None:
    with pytest.raises(ParseError, match="UNSUPPORTED_FILE_TYPE"):
        validate_upload_type("legacy.xls", "application/vnd.ms-excel")


def test_s3_checksum_uses_the_required_base64_header_value() -> None:
    assert (
        ObjectStorageClient._checksum_header("00" * 32)
        == "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA="
    )


def test_s3_verification_requests_provider_checksum_metadata() -> None:
    class FakeClient:
        def __init__(self) -> None:
            self.kwargs: dict[str, object] = {}

        def head_object(self, **kwargs: object) -> dict[str, object]:
            self.kwargs = kwargs
            return {
                "ContentLength": 1,
                "ContentType": "text/plain",
                "ChecksumSHA256": ObjectStorageClient._checksum_header("00" * 32),
            }

    client = ObjectStorageClient(Settings())
    fake = FakeClient()
    client._client = fake
    client.verify("object", "00" * 32, 1, "text/plain")
    assert fake.kwargs["ChecksumMode"] == "ENABLED"


def test_worker_reconciles_due_retries_and_expired_leases() -> None:
    worker = (Path(__file__).parents[2] / "src" / "ima" / "workers" / "main.py").read_text()
    assert "status IN ('queued','retryable')" in worker
    assert "lease_expires_at < :now" in worker
    assert "storage_cleanup_jobs" in worker
    assert "defer_async(" in worker


def test_stage_rows_begin_blocked_until_the_predecessor_succeeds() -> None:
    root = Path(__file__).parents[2]
    migration_path = root / "migrations" / "versions" / "20260825_0006_storage_ingestion.py"
    migration = migration_path.read_text()
    service = (root / "src" / "ima" / "infrastructure" / "tasks" / "service.py").read_text()
    assert "'blocked','queued','running'" in migration
    assert ":stage,'blocked'" in service


def test_completion_persists_stage_rows_with_the_published_version_before_delivery() -> None:
    storage = (Path(__file__).parents[2] / "src" / "ima" / "application" / "storage.py").read_text()
    publish = storage.index("UPDATE ima.documents SET current_version")
    stage_rows = storage.index("await self.jobs.create_ingestion_jobs")
    delivery = storage.index("await self.jobs.defer_ingestion_parse")
    transaction_end = storage.rindex("        await self.jobs.defer_ingestion_parse")
    assert publish < stage_rows < transaction_end <= delivery


def test_retry_cancel_and_status_are_scoped_to_the_current_file_version() -> None:
    storage = (Path(__file__).parents[2] / "src" / "ima" / "application" / "storage.py").read_text()
    retry = storage.index("async def retry")
    cancel = storage.index("async def cancel")
    status = storage.index("async def status")
    retry_section = storage[retry:cancel]
    cancel_section = storage[cancel:status]
    status_section = storage[status:]
    assert "document_id=:id AND version=:version" in retry_section
    assert '"version": row["version"]' in retry_section
    assert "document_id=:id AND version=:version" in cancel_section
    assert "document_id=:id AND version=:version" in status_section


def test_parser_rejects_invalid_json() -> None:
    with pytest.raises(ParseError, match="MALFORMED_DOCUMENT"):
        Parser(Settings()).parse(b"{", "bad.json", "application/json")


def test_markdown_heading_paths_are_prefixed_to_chunks() -> None:
    parsed = Parser(Settings()).parse(
        b"# Overall\n\nintro\n\n## Details\n\nbody", "note.md", "text/markdown"
    )
    chunks = structure_aware_chunks(parsed.blocks, 100, 0, 10)
    assert chunks[0].heading_path == ("Overall",)
    assert chunks[0].text_content.startswith("Overall\n\n")
    assert chunks[0].chunker_version == "structure-aware-v1"
    assert chunks[0].metadata["chunker_config_digest"]
    assert chunks[1].heading_path == ("Overall", "Details")
    assert chunks[1].text_content.startswith("Overall > Details\n\n")


def test_markdown_table_stays_atomic_and_repeats_header_when_split() -> None:
    parsed = Parser(Settings()).parse(
        b"# Data\n\n| name | value |\n| --- | --- |\n| alpha | one |\n| beta | two |",
        "table.md",
        "text/markdown",
    )
    whole = structure_aware_chunks(parsed.blocks, 100, 0, 10)
    assert len(whole) == 1
    assert whole[0].chunk_type == "table"
    split = structure_aware_chunks(parsed.blocks, 8, 0, 10)
    assert len(split) > 1
    assert all("name | value" in chunk.text_content for chunk in split)
    assert all(chunk.metadata["table_header_repeated"] for chunk in split)


def test_table_without_explicit_header_keeps_data_rows_once() -> None:
    parsed = Parser(Settings()).parse(
        b"first | value\nsecond | value\nthird | value",
        "data.md",
        "text/markdown",
    )
    table = next(block for block in parsed.blocks if block.kind == "table")
    assert table.has_header is False
    chunks = structure_aware_chunks(parsed.blocks, 5, 0, 10)
    assert all("first | value" not in chunk.text_content for chunk in chunks[1:])
    assert sum(chunk.text_content.count("first | value") for chunk in chunks) == 1
    assert all(chunk.text_content.strip() for chunk in chunks)


def test_english_sentence_boundaries_avoid_decimal_abbreviation_and_url_dots() -> None:
    blocks = (
        ParsedBlock(
            "prose",
            "One two three. Four five six. Dr. Smith arrived. 3.14 is pi. "
            "Visit https://example.com now.",
        ),
    )
    chunks = structure_aware_chunks(blocks, 12, 0, 10)
    rendered = "\n".join(chunk.text_content for chunk in chunks)
    assert "One two three." in rendered
    assert "Four five six." in rendered
    assert "Dr. Smith arrived." in rendered
    assert "3.14 is pi." in rendered
    assert "https://example.com now." in rendered
    assert "Dr.\nSmith" not in rendered


def test_heading_prefix_is_included_in_the_final_token_budget() -> None:
    blocks = (
        ParsedBlock("heading", "Data", ("Data",), heading_level=1),
        ParsedBlock("prose", "abcdefghijklmnopqrst", ("Data",)),
    )
    chunks = structure_aware_chunks(blocks, 6, 0, 10)
    assert chunks
    assert all(chunk.token_count <= 6 for chunk in chunks)


def test_empty_ir_and_heading_only_documents_produce_no_chunks() -> None:
    assert structure_aware_chunks((), 100, 0, 10) == ()
    assert (
        structure_aware_chunks(
            (ParsedBlock("heading", "Only heading", ("Only heading",), heading_level=1),),
            100,
            0,
            10,
        )
        == ()
    )


def test_docx_blocks_keep_paragraph_table_order_and_heading_level() -> None:
    document = Document()
    document.add_paragraph("Section", style="Heading 2")
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "1"
    document.add_paragraph("After table")
    stream = BytesIO()
    document.save(stream)

    parsed = Parser(Settings()).parse(
        stream.getvalue(),
        "ordered.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert [block.kind for block in parsed.blocks] == ["heading", "prose", "table", "prose"]
    assert parsed.blocks[0].heading_level == 2
    assert parsed.blocks[2].metadata["header"] is None


def test_pdf_blocks_keep_page_metadata(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        Parser,
        "_pdf",
        lambda _self, _data: (
            ParsedBlock("prose", "page one", page_start=1, page_end=1),
            ParsedBlock("prose", "page two", page_start=2, page_end=2),
        ),
    )
    parsed = Parser(Settings()).parse(b"pdf", "sample.pdf", "application/pdf")
    chunks = structure_aware_chunks(parsed.blocks, 100, 0, 10)
    assert [(chunk.page_start, chunk.page_end) for chunk in chunks] == [(1, 2)]


def test_xlsx_chunks_keep_sheet_and_header_metadata() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Name", "Amount"])
    for index in range(1, 8):
        sheet.append([f"A{index}", index])
    stream = BytesIO()
    workbook.save(stream)

    parsed = Parser(Settings()).parse(
        stream.getvalue(),
        "sales.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    table = next(block for block in parsed.blocks if block.kind == "table")
    assert table.sheet_name == "Sales"
    assert table.row_start == 1
    assert table.row_end == 8
    assert table.has_header is True
    chunks = structure_aware_chunks(parsed.blocks, 8, 0, 10)
    assert len(chunks) > 1
    assert all(chunk.sheet_name == "Sales" for chunk in chunks)
    assert all("Name | Amount" in chunk.text_content for chunk in chunks)


def test_no_header_table_does_not_fill_header_metadata() -> None:
    parsed = Parser(Settings()).parse(
        b"first | value\nsecond | value",
        "data.md",
        "text/markdown",
    )
    table = next(block for block in parsed.blocks if block.kind == "table")
    assert table.has_header is False
    assert table.metadata["header"] is None


def test_xlsx_without_a_header_signal_keeps_header_metadata_empty() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B"])
    sheet.append(["C", "D"])
    stream = BytesIO()
    workbook.save(stream)

    parsed = Parser(Settings()).parse(
        stream.getvalue(),
        "strings.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    table = next(block for block in parsed.blocks if block.kind == "table")
    assert table.has_header is False
    assert table.metadata["header"] is None


def test_common_abbreviations_survive_sentence_and_character_fallbacks() -> None:
    text = "Inc. Closed later. Ltd. remains. e.g. examples continue. " + ("details " * 80)
    chunks = structure_aware_chunks((ParsedBlock("prose", text),), 8, 0, 100)
    rendered = "\n".join(chunk.text_content for chunk in chunks)
    assert "Inc.\nClosed" not in rendered
    assert "Ltd.\nremains" not in rendered
    assert "e.\n.g." not in rendered


def test_oversized_heading_path_is_truncated_without_dropping_body() -> None:
    blocks = (
        ParsedBlock(
            "prose",
            "Body survives an oversized heading path.",
            ("A very long heading path that exceeds the token budget", "Leaf"),
        ),
    )
    chunks = structure_aware_chunks(blocks, 5, 0, 20)
    assert chunks
    assert "Body" in "\n".join(chunk.text_content for chunk in chunks)
    assert all(chunk.token_count <= 5 for chunk in chunks)


def test_token_estimator_counts_cjk_as_one_and_latin_as_quarter() -> None:
    assert estimate_tokens("中文中文abcd") == 5


def test_oversized_atomic_block_falls_back_to_character_chunks() -> None:
    blocks = (ParsedBlock("code", "abcdefghij"),)
    chunks = structure_aware_chunks(blocks, 2, 0, 10)
    assert len(chunks) > 1
    assert all(chunk.token_count <= 2 for chunk in chunks)
    assert all(chunk.metadata["oversized_fallback"] for chunk in chunks)


def test_embed_stage_resolves_through_the_shared_embedding_target() -> None:
    # Ingestion must not query the assignment table directly: resolution comes
    # from the scene-default embedding target retrieval already honours.
    source = (
        Path(__file__).parents[2] / "src" / "ima" / "infrastructure" / "tasks" / "ingestion.py"
    ).read_text()
    embed = source[source.index('name="ima.ingestion.embed"') :]
    embed = embed[: embed.index('name="ima.ingestion.cleanup"')]
    assert "kb_profile_assignments" not in embed
    assert "embedding_target(" in embed
    assert '"NO_ASSIGNMENT"' in embed
    assert "managed_embeddings(" in embed


def test_chunk_pipeline_rejects_empty_sets_and_keeps_generation_history() -> None:
    root = Path(__file__).parents[2]
    source = (root / "src" / "ima" / "infrastructure" / "tasks" / "ingestion.py").read_text()
    assert 'if not records:\n                    await failed(conn, job, "NO_CHUNKS"' in source
    assert 'if not chunks:\n                    await failed(conn, job, "NO_CHUNKS"' in source
    assert "INSERT INTO ima.document_derived_text" in source
    assert "UPDATE ima.document_file_versions SET generation=:generation" in source


def test_structure_chunk_constraints_are_added_without_initial_table_scan() -> None:
    migration = (
        Path(__file__).parents[2]
        / "migrations"
        / "versions"
        / "20260914_0015_structure_aware_chunks.py"
    ).read_text()
    assert "NOT VALID" in migration
    assert "VALIDATE CONSTRAINT ck_document_chunks_metadata_object" in migration
    assert "VALIDATE CONSTRAINT ck_document_chunks_type" in migration


def test_embedding_batches_persist_partial_progress_for_retry() -> None:
    source = (
        Path(__file__).parents[2] / "src" / "ima" / "infrastructure" / "tasks" / "ingestion.py"
    ).read_text()
    embed = source[source.index('name="ima.ingestion.embed"') :]
    embed = embed[: embed.index('name="ima.ingestion.cleanup"')]
    assert 'pending_chunks = [row for row in chunks if row["embedding_status"] != "ready"]' in embed
    assert "completed = len(chunks) - len(pending_chunks)" in embed
    assert "completed += len(batch)" in embed
    assert "except GatewayError as exc" in embed
    assert "await failed(conn, job, exc.code, retryable=True)" in embed
